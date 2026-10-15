"""Excel helpers shared by all bulk imports and exports."""

from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.config import settings

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class ImportFileError(ValueError):
    pass


def _norm(header: Any) -> str:
    return str(header or "").strip().lower().replace("*", "").replace(" ", "_")


def build_template(columns: list[tuple[str, bool, str]], example: list[Any] | None = None) -> bytes:
    """columns: (name, required, help text). Required headers get a '*'."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    for i, (name, required, _help) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=i, value=f"{name}*" if required else name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F6F5C" if required else "5B6B66")
        ws.column_dimensions[get_column_letter(i)].width = max(14, len(name) + 4)
    if example:
        for i, v in enumerate(example, start=1):
            ws.cell(row=2, column=i, value=v)
    ws.freeze_panes = "A2"

    guide = wb.create_sheet("Instructions")
    guide.append(["Column", "Required", "How to fill"])
    for c in guide[1]:
        c.font = Font(bold=True)
    for name, required, help_text in columns:
        guide.append([name, "Yes" if required else "No", help_text])
    guide.column_dimensions["A"].width = 22
    guide.column_dimensions["C"].width = 80

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def read_rows(content: bytes, required: list[str]) -> list[tuple[int, dict[str, Any]]]:
    """Return (excel_row_number, {column: value}) for every non-empty row of the first sheet."""
    try:
        wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - any parse failure means a bad file
        raise ImportFileError("Could not read the file. Upload an .xlsx file made from the template.") from exc
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    try:
        headers = [_norm(h) for h in next(rows)]
    except StopIteration:
        raise ImportFileError("The file is empty") from None

    missing = [c for c in required if c not in headers]
    if missing:
        raise ImportFileError(f"Missing required columns: {', '.join(missing)}")

    out: list[tuple[int, dict[str, Any]]] = []
    for idx, values in enumerate(rows, start=2):
        if values is None or all(v is None or str(v).strip() == "" for v in values):
            continue
        row = {}
        for h, v in zip(headers, values, strict=False):
            if not h:
                continue
            if isinstance(v, str):
                v = v.strip() or None
            row[h] = v
        out.append((idx, row))
        if len(out) > settings.MAX_UPLOAD_ROWS:
            raise ImportFileError(f"Too many rows; the limit is {settings.MAX_UPLOAD_ROWS} per file")
    wb.close()
    return out


def as_text(v: Any) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s or None


def as_bool(v: Any) -> bool:
    return str(v).strip().lower() in {"yes", "y", "true", "1"} if v is not None else False


def as_list(v: Any) -> list[str]:
    s = as_text(v)
    return [p.strip() for p in s.split(",") if p.strip()] if s else []


def export_rows(headers: list[str], rows: list[list[Any]], sheet: str = "Export") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, len(h) + 4)
    ws.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def pydantic_errors(exc: Exception) -> list[str]:
    errors = getattr(exc, "errors", None)
    if not callable(errors):
        return [str(exc)]
    out = []
    for e in errors():
        loc = ".".join(str(p) for p in e.get("loc", []) if p != "__root__")
        msg = e.get("msg", "invalid").removeprefix("Value error, ")
        out.append(f"{loc}: {msg}" if loc else msg)
    return out
