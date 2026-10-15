"""Orders API (M6). Business rules live in app.services.orders; this module handles access and I/O."""

import uuid
from collections import defaultdict
from datetime import date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO

from fastapi import APIRouter, Depends, File, HTTPException, Query, Response, UploadFile, status
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session, selectinload

from app.api.v1.customers import get_customer_or_404
from app.core.database import get_db
from app.core.deps import ALL_STAFF, MANAGERS, require_roles
from app.core.validators import normalize_phone
from app.models import Customer, Message, Order, OrderEvent, OrderLine, OrderSource, OrderStatus, Product, Role, User
from app.schemas.common import ImportResult, ImportRowError, Page
from app.schemas.orders import (
    CancelIn,
    CreditInfo,
    DispatchIn,
    InvoiceIn,
    OrderCreate,
    OrderDetail,
    OrderEventOut,
    OrderLinesUpdate,
    OrderSummary,
    RemarkIn,
    SourceMessage,
)
from app.services import audit, excel
from app.services.ai.matcher import norm
from app.services import orders as svc
from app.services.orders import IST, OrderError

router = APIRouter(prefix="/orders", tags=["Orders"])
staff = require_roles(*ALL_STAFF)
writers = require_roles(Role.owner, Role.admin, Role.sales)
managers = require_roles(*MANAGERS)
invoicers = require_roles(Role.owner, Role.admin, Role.accounts)

S = OrderStatus
SALES_CAN_CANCEL = {S.draft, S.quoted, S.expired, S.confirmed, S.on_hold}


def _err(exc: OrderError) -> HTTPException:
    return HTTPException(exc.http, exc.message)


def scoped(stmt, user: User):
    """Tenant filter; Sales sees only orders of customers assigned to them (FR-M1-05)."""
    stmt = stmt.where(Order.tenant_id == user.tenant_id)
    if user.role == Role.sales:
        stmt = stmt.where(Order.customer_id.in_(
            select(Customer.id).where(Customer.tenant_id == user.tenant_id, Customer.assigned_user_id == user.id)))
    return stmt


def get_order_or_404(db: Session, user: User, order_id: uuid.UUID, lock: bool = False) -> Order:
    stmt = scoped(select(Order).where(Order.id == order_id), user)
    if lock:
        stmt = stmt.with_for_update()
    order = db.scalar(stmt)
    if order is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Order not found")
    return order


def allowed_actions(user: User, order: Order) -> list[str]:
    st, role = order.status, user.role
    mgr = role in MANAGERS
    acts: list[str] = []
    if role in (Role.owner, Role.admin, Role.sales):
        if st in svc.EDITABLE:
            acts.append("edit")
        if st in (S.draft, S.quoted, S.expired):
            acts.append("send_quote")
        if st in (S.draft, S.quoted):
            acts.append("confirm")
        if st in SALES_CAN_CANCEL or (mgr and st == S.approved):
            acts.append("cancel")
    if mgr and st in (S.confirmed, S.on_hold):
        acts.append("approve")
    if mgr and st == S.on_hold:
        acts.append("reject")
    if role in (Role.owner, Role.admin, Role.accounts) and st == S.approved:
        acts.append("invoice")
    if mgr and st == S.invoiced:
        acts.append("dispatch")
    if mgr and st == S.dispatched:
        acts.append("deliver")
    return acts


def _names(db: Session, tenant_id: uuid.UUID) -> dict[uuid.UUID, str]:
    return {u.id: u.name for u in db.scalars(select(User).where(User.tenant_id == tenant_id))}


def _detail(db: Session, user: User, order: Order) -> OrderDetail:
    db.refresh(order)
    customer = db.get(Customer, order.customer_id)
    names = _names(db, user.tenant_id)
    events = db.scalars(select(OrderEvent).where(OrderEvent.order_id == order.id)
                        .order_by(OrderEvent.created_at, OrderEvent.id)).all()
    limit = Decimal(customer.credit_limit or 0)
    outstanding = svc.credit_exposure(db, customer)
    data = {c: getattr(order, c) for c in OrderDetail.model_fields
            if c not in {"customer_name", "customer_whatsapp", "sales_user_name", "created_by_name", "events",
                         "credit", "allowed_actions", "lines", "source_message"}}
    src = db.get(Message, order.source_message_id) if order.source_message_id else None
    return OrderDetail(
        **data,
        lines=order.lines,
        customer_name=customer.business_name,
        customer_whatsapp=customer.whatsapp_number,
        sales_user_name=names.get(order.sales_user_id),
        created_by_name=names.get(order.created_by_user_id),
        events=[OrderEventOut.model_validate(e).model_copy(update={"user_name": names.get(e.user_id)}) for e in events],
        credit=CreditInfo(credit_limit=limit, outstanding=outstanding,
                          available=(limit - outstanding) if limit > 0 else None),
        allowed_actions=allowed_actions(user, order),
        source_message=SourceMessage(
            id=src.id, msg_type=src.msg_type.value, text=src.text, transcript=(src.payload or {}).get("transcript"),
            has_media=bool(src.media_path), media_mime=src.media_mime, media_filename=src.media_filename,
            created_at=src.created_at) if src else None,
    )


def _items(db: Session, tenant_id: uuid.UUID, lines) -> list[dict]:
    products = svc.load_products(db, tenant_id, [l.product_id for l in lines])
    return [{"product": products[l.product_id], "quantity": l.quantity, "raw_text": l.raw_text} for l in lines]


def _filtered(user: User, *, status_: list[OrderStatus] | None, date_from: date | None, date_to: date | None,
              customer_id: uuid.UUID | None, sales_user_id: uuid.UUID | None, source: OrderSource | None,
              q: str | None):
    stmt = scoped(select(Order).join(Customer, Customer.id == Order.customer_id), user)
    if status_:
        stmt = stmt.where(Order.status.in_(status_))
    if date_from:
        stmt = stmt.where(Order.created_at >= datetime.combine(date_from, time.min, IST))
    if date_to:
        stmt = stmt.where(Order.created_at < datetime.combine(date_to + timedelta(days=1), time.min, IST))
    if customer_id:
        stmt = stmt.where(Order.customer_id == customer_id)
    if sales_user_id:
        stmt = stmt.where(Order.sales_user_id == sales_user_id)
    if source:
        stmt = stmt.where(Order.source == source)
    if q:
        like = f"%{q.strip()}%"
        stmt = stmt.where(or_(Order.number.ilike(like), Customer.business_name.ilike(like),
                              Customer.whatsapp_number.ilike(like)))
    return stmt


@router.get("", response_model=Page[OrderSummary])
def list_orders(
    status_filter: list[OrderStatus] | None = Query(default=None, alias="status"),
    date_from: date | None = None,
    date_to: date | None = None,
    customer_id: uuid.UUID | None = None,
    sales_user_id: uuid.UUID | None = None,
    source: OrderSource | None = None,
    q: str | None = None,
    page: int = Query(default=1, ge=1),
    size: int = Query(default=25, ge=1, le=200),
    db: Session = Depends(get_db),
    user: User = Depends(staff),
):
    """FR-M6-01. The approval queue is status=on_hold."""
    stmt = _filtered(user, status_=status_filter, date_from=date_from, date_to=date_to, customer_id=customer_id,
                     sales_user_id=sales_user_id, source=source, q=q)
    total = db.scalar(select(func.count()).select_from(stmt.subquery()))
    orders = db.scalars(stmt.order_by(Order.created_at.desc()).offset((page - 1) * size).limit(size)).all()
    names = _names(db, user.tenant_id)
    customers = {c.id: c.business_name for c in db.scalars(
        select(Customer).where(Customer.id.in_({o.customer_id for o in orders})))} if orders else {}
    counts = dict(db.execute(select(OrderLine.order_id, func.count()).where(
        OrderLine.order_id.in_([o.id for o in orders])).group_by(OrderLine.order_id)).all()) if orders else {}
    items = [OrderSummary(
        id=o.id, number=o.number, customer_id=o.customer_id, customer_name=customers.get(o.customer_id, ""),
        sales_user_id=o.sales_user_id, sales_user_name=names.get(o.sales_user_id), source=o.source,
        status=o.status, grand_total=o.grand_total, line_count=counts.get(o.id, 0), hold_reasons=o.hold_reasons,
        quote_valid_until=o.quote_valid_until, created_at=o.created_at, updated_at=o.updated_at,
    ) for o in orders]
    return Page(items=items, total=total, page=page, size=size)


@router.get("/counts", response_model=dict[str, int])
def order_counts(db: Session = Depends(get_db), user: User = Depends(staff)):
    rows = db.execute(scoped(select(Order.status, func.count()), user).group_by(Order.status)).all()
    out = {s.value: 0 for s in OrderStatus}
    out.update({s.value: n for s, n in rows})
    return out


@router.post("", response_model=OrderDetail, status_code=status.HTTP_201_CREATED)
def create_order(body: OrderCreate, db: Session = Depends(get_db), user: User = Depends(writers)):
    """FR-M6-02: staff create an order for a customer; prices follow M3 rules."""
    customer = get_customer_or_404(db, user, body.customer_id)
    try:
        order = svc.create_order(db, user, customer, _items(db, user.tenant_id, body.lines), notes=body.notes)
        if body.action == "quote":
            svc.send_quote(db, user, order)
        elif body.action == "confirm":
            svc.confirm(db, order, user=user)
    except OrderError as exc:
        db.rollback()
        raise _err(exc) from None
    svc.commit_and_send(db)
    return _detail(db, user, order)


# ---------- Excel (declared before /{order_id}) ----------
IMPORT_COLUMNS = [
    ("order_ref", True, "Your reference that groups rows into one order, e.g. A1. Rows with the same ref become one order"),
    ("customer_whatsapp", True, "Customer's WhatsApp number as saved in Customers, e.g. 9876543210"),
    ("sku", True, "Product SKU code"),
    ("quantity", True, "Quantity, e.g. 10"),
    ("notes", False, "Optional note for the order (taken from the first row of each order)"),
]


@router.get("/import/template")
def order_import_template(_: User = Depends(writers)):
    content = excel.build_template(IMPORT_COLUMNS, ["A1", "9876543210", "SALT-1KG", 10, "Deliver by Friday"])
    return Response(content, media_type=excel.XLSX_MIME,
                    headers={"Content-Disposition": 'attachment; filename="order_import_template.xlsx"'})


@router.post("/import", response_model=ImportResult)
async def import_orders(file: UploadFile = File(...), db: Session = Depends(get_db), user: User = Depends(writers)):
    """FR-M6-03: creates Draft orders. An order with any bad row is skipped as a whole."""
    try:
        rows = excel.read_rows(await file.read(), ["order_ref", "customer_whatsapp", "sku", "quantity"])
    except excel.ImportFileError as exc:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, str(exc)) from None

    tid = user.tenant_id
    products = {p.sku.lower(): p for p in db.scalars(select(Product).where(Product.tenant_id == tid))}
    groups: dict[str, list[tuple[int, dict]]] = defaultdict(list)
    errors: list[ImportRowError] = []
    for row_no, r in rows:
        ref = excel.as_text(r.get("order_ref"))
        if not ref:
            errors.append(ImportRowError(row=row_no, errors=["order_ref: required"]))
            continue
        groups[ref].append((row_no, r))

    created = 0
    for ref, group in groups.items():
        group_errors: list[ImportRowError] = []
        numbers = {excel.as_text(r.get("customer_whatsapp")) for _, r in group}
        customer = None
        if len(numbers) > 1:
            group_errors.append(ImportRowError(row=group[0][0], errors=[
                f"order_ref {ref}: all rows of one order must have the same customer_whatsapp"]))
        else:
            raw = numbers.pop()
            try:
                number = normalize_phone(raw or "")
                stmt = select(Customer).where(Customer.tenant_id == tid, Customer.whatsapp_number == number)
                if user.role == Role.sales:
                    stmt = stmt.where(Customer.assigned_user_id == user.id)
                customer = db.scalar(stmt)
                if customer is None:
                    group_errors.append(ImportRowError(row=group[0][0], errors=[
                        f"customer_whatsapp: no customer with number {raw}"]))
            except ValueError:
                group_errors.append(ImportRowError(row=group[0][0], errors=[f"customer_whatsapp: '{raw}' is not valid"]))
        items: list[dict] = []
        for row_no, r in group:
            errs = []
            sku = (excel.as_text(r.get("sku")) or "").lower()
            product = products.get(sku)
            if product is None:
                errs.append(f"sku: '{r.get('sku')}' not found")
            elif not product.is_active:
                errs.append(f"sku: '{product.sku}' is inactive")
            try:
                qty = Decimal(str(r.get("quantity")).replace(",", ""))
                if qty <= 0:
                    raise InvalidOperation
            except (InvalidOperation, ValueError):
                errs.append("quantity: must be a number above zero")
                qty = None
            if errs:
                group_errors.append(ImportRowError(row=row_no, errors=errs))
            elif product is not None and qty is not None:
                items.append({"product": product, "quantity": qty})
        if group_errors or customer is None:
            errors.extend(group_errors)
            continue
        try:
            with db.begin_nested():
                svc.create_order(db, user, customer, items, source=OrderSource.excel,
                                 notes=excel.as_text(group[0][1].get("notes")))
            created += 1
        except OrderError as exc:
            errors.append(ImportRowError(row=group[0][0], errors=[f"order_ref {ref}: {exc.message}"]))
    db.commit()
    return ImportResult(total_rows=len(rows), created=created, updated=0, failed=len(errors), errors=errors)


@router.get("/export")
def export_orders(
    status_filter: list[OrderStatus] | None = Query(default=None, alias="status"),
    date_from: date | None = None,
    date_to: date | None = None,
    customer_id: uuid.UUID | None = None,
    sales_user_id: uuid.UUID | None = None,
    source: OrderSource | None = None,
    q: str | None = None,
    db: Session = Depends(get_db),
    user: User = Depends(staff),
):
    """FR-M6-11: orders and order lines, with the same filters as the list."""
    stmt = _filtered(user, status_=status_filter, date_from=date_from, date_to=date_to, customer_id=customer_id,
                     sales_user_id=sales_user_id, source=source, q=q)
    orders = db.scalars(stmt.options(selectinload(Order.lines)).order_by(Order.created_at.desc())).all()
    names = _names(db, user.tenant_id)
    customers = {c.id: c for c in db.scalars(select(Customer).where(Customer.tenant_id == user.tenant_id))}

    wb = Workbook()
    ws = wb.active
    ws.title = "Orders"
    oh = ["order_number", "date", "customer", "customer_whatsapp", "sales_user", "source", "status",
          "taxable_value", "gst", "total", "invoice_number", "hold_reasons", "cancel_reason", "notes"]
    ws.append(oh)
    lines_ws = wb.create_sheet("Order lines")
    lh = ["order_number", "sku", "product", "uom", "quantity", "free_qty", "rate", "discount", "scheme",
          "taxable_value", "gst_rate", "gst", "line_total", "invoiced_qty", "short_supply"]
    lines_ws.append(lh)
    for o in orders:
        c = customers.get(o.customer_id)
        ws.append([o.number, o.created_at.astimezone(IST).strftime("%Y-%m-%d %H:%M"),
                   c.business_name if c else "", c.whatsapp_number if c else "", names.get(o.sales_user_id),
                   o.source.value, svc.label(o.status), float(o.taxable_total), float(o.gst_total),
                   float(o.grand_total), o.invoice_number, o.hold_reasons, o.cancel_reason, o.notes])
        for l in o.lines:
            short = float(l.quantity - l.invoiced_qty) if l.invoiced_qty is not None else None
            lines_ws.append([o.number, l.sku, l.name, l.uom, float(l.quantity), float(l.free_qty),
                             float(l.unit_price), float(l.discount), l.scheme_name, float(l.taxable_value),
                             float(l.gst_rate), float(l.gst_amount), float(l.line_total),
                             float(l.invoiced_qty) if l.invoiced_qty is not None else None, short])
    for sheet, headers in ((ws, oh), (lines_ws, lh)):
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for i, h in enumerate(headers, start=1):
            sheet.column_dimensions[get_column_letter(i)].width = max(12, len(h) + 4)
        sheet.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type=excel.XLSX_MIME,
                    headers={"Content-Disposition": 'attachment; filename="orders.xlsx"'})


# ---------- single order ----------
@router.get("/{order_id}", response_model=OrderDetail)
def get_order(order_id: uuid.UUID, db: Session = Depends(get_db), user: User = Depends(staff)):
    return _detail(db, user, get_order_or_404(db, user, order_id))


def _act(db: Session, user: User, order_id: uuid.UUID, fn) -> OrderDetail:
    order = get_order_or_404(db, user, order_id, lock=True)
    try:
        fn(order)
    except OrderError as exc:
        if exc.code == "quote_expired":  # the expiry itself is real and must be kept
            svc.commit_and_send(db)
        else:
            db.rollback()
        raise _err(exc) from None
    svc.commit_and_send(db)
    return _detail(db, user, order)


@router.put("/{order_id}/lines", response_model=OrderDetail)
def edit_order_lines(order_id: uuid.UUID, body: OrderLinesUpdate, db: Session = Depends(get_db),
                     user: User = Depends(writers)):
    return _act(db, user, order_id,
                lambda o: svc.edit_lines(db, user, o, _items(db, user.tenant_id, body.lines), body.notes))


@router.post("/{order_id}/send-quote", response_model=OrderDetail)
def send_quote(order_id: uuid.UUID, db: Session = Depends(get_db), user: User = Depends(writers)):
    def run(o: Order):
        if o.status not in (S.draft, S.quoted, S.expired):
            raise OrderError(f"Cannot send a quote for a {svc.label(o.status)} order", "invalid_transition", 409)
        svc.send_quote(db, user, o, revised=o.revision > 1)
    return _act(db, user, order_id, run)


@router.post("/{order_id}/confirm", response_model=OrderDetail)
def confirm(order_id: uuid.UUID, db: Session = Depends(get_db), user: User = Depends(writers)):
    """Staff confirm on the buyer's instruction (phone call, visit). Credit check runs here."""
    return _act(db, user, order_id, lambda o: svc.confirm(db, o, user=user))


@router.post("/{order_id}/approve", response_model=OrderDetail)
def approve(order_id: uuid.UUID, body: RemarkIn, db: Session = Depends(get_db), user: User = Depends(managers)):
    return _act(db, user, order_id, lambda o: svc.approve(db, user, o, body.remark))


@router.post("/{order_id}/reject", response_model=OrderDetail)
def reject(order_id: uuid.UUID, body: RemarkIn, db: Session = Depends(get_db), user: User = Depends(managers)):
    return _act(db, user, order_id, lambda o: svc.reject(db, user, o, body.remark or ""))


@router.post("/{order_id}/cancel", response_model=OrderDetail)
def cancel(order_id: uuid.UUID, body: CancelIn, db: Session = Depends(get_db), user: User = Depends(writers)):
    def run(o: Order):
        if user.role == Role.sales and o.status == S.approved:
            raise OrderError("Only Owner or Admin can cancel an approved order", "forbidden", 403)
        svc.cancel(db, o, body.reason, user=user)
    return _act(db, user, order_id, run)


@router.post("/{order_id}/invoice", response_model=OrderDetail)
def mark_invoiced(order_id: uuid.UUID, body: InvoiceIn, db: Session = Depends(get_db),
                  user: User = Depends(invoicers)):
    qty = {l.line_id: l.invoiced_qty for l in body.lines}
    return _act(db, user, order_id, lambda o: svc.mark_invoiced(db, user, o, body.invoice_number, body.invoice_date, qty))


@router.post("/{order_id}/dispatch", response_model=OrderDetail)
def dispatch(order_id: uuid.UUID, body: DispatchIn, db: Session = Depends(get_db), user: User = Depends(managers)):
    return _act(db, user, order_id,
                lambda o: svc.dispatch(db, user, o, body.vehicle_number, body.transporter, body.lr_number))


@router.post("/{order_id}/deliver", response_model=OrderDetail)
def deliver(order_id: uuid.UUID, db: Session = Depends(get_db), user: User = Depends(managers)):
    return _act(db, user, order_id, lambda o: svc.deliver(db, user, o))


@router.post("/{order_id}/lines/{line_id}/save-alias", response_model=OrderDetail)
def save_alias(order_id: uuid.UUID, line_id: uuid.UUID, db: Session = Depends(get_db), user: User = Depends(writers)):
    """FR-M5-09: keep the buyer's wording as a product alias so it matches next time."""
    order = get_order_or_404(db, user, order_id)
    line = next((l for l in order.lines if l.id == line_id), None)
    if line is None or not line.raw_text:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "This line has no buyer wording to save")
    product = db.scalar(select(Product).where(Product.id == line.product_id, Product.tenant_id == user.tenant_id)
                        .with_for_update())
    alias = norm(line.raw_text)[:150]
    existing = {norm(a) for a in (product.aliases or [])} | {norm(product.name), norm(product.sku)}
    if alias and alias not in existing:
        old = audit.snapshot(product)
        product.aliases = [*(product.aliases or []), alias]
        audit.record(db, user, "product", product.id, "update", old, audit.snapshot(product))
    db.commit()
    return _detail(db, user, order)
