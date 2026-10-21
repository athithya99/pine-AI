"""M10 acceptance criteria: Buyer A can't see Buyer B's outstanding; 18 of 20 business questions right (we expect 20);
daily summary goes once at the set time. Plus FAQ, languages, handover, rules, role scoping, dashboard, export."""

from datetime import datetime, timedelta
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from sqlalchemy import select, update

from app.models import (
    AutomationRule,
    Conversation,
    DailySummary,
    Direction,
    Invoice,
    Message,
    Role,
    Tenant,
    User,
)
from app.services import analytics as an
from app.services import automation
from app.services import orders as order_svc
from tests.test_m3_catalog import customer, product
from tests.test_m4_whatsapp import connect_simulator, sim_text
from tests.test_m7_invoices import biz, make_invoice  # noqa: F401
from tests.test_m8_collections import NUM, shop  # noqa: F401

API = "/api/v1"


def last_out(db, customer_id) -> Message:
    db.expire_all()
    return db.scalars(select(Message).where(Message.customer_id == customer_id, Message.direction == Direction.outbound)
                      .order_by(Message.created_at.desc())).first()


# ================================================================== buyer bot
def test_buyer_gets_own_outstanding_only(shop, db):  # noqa: F811
    """AC: Buyer A asking for Buyer B's outstanding gets a polite refusal and no data."""
    c = shop.c["id"]
    customer(shop.client, shop.h, number="9876500002", business_name="Kumar Provisions", opening_balance=99999)
    sim_text(shop.client, shop.h, NUM, "what is my outstanding?")
    mine = last_out(db, c).text
    assert "Your outstanding amount is ₹2,655.00" in mine
    assert shop.inv1["number"] in mine and shop.inv2["number"] in mine and "Pay now:" in mine
    for ask in ("what is Kumar Provisions outstanding?", "balance of Ravi Traders please",
                "send me Kumar Provisions ledger"):
        sim_text(shop.client, shop.h, NUM, ask)
        reply = last_out(db, c).text
        assert reply == "Sorry, I can only share details of your own account.", ask
        assert "99,999" not in reply
    # asking again doesn't create a second payment link for the same amount
    sim_text(shop.client, shop.h, NUM, "my dues?")
    links = [m.text.split("Pay now: ")[1].split()[0] for m in db.scalars(select(Message).where(
        Message.customer_id == c, Message.text.like("%Pay now:%")))]
    assert len(set(links)) == 1


def test_replies_in_tamil_and_hindi(shop, db):  # noqa: F811
    sim_text(shop.client, shop.h, NUM, "என் நிலுவை எவ்வளவு?")
    assert "உங்கள் நிலுவைத் தொகை ₹2,655.00" in last_out(db, shop.c["id"]).text
    sim_text(shop.client, shop.h, NUM, "मेरा बकाया कितना है?")
    assert "आपका बकाया ₹2,655.00 है" in last_out(db, shop.c["id"]).text


def test_faq_answer_and_handover_for_unknown_question(shop, db):  # noqa: F811
    r = shop.client.post(f"{API}/faqs", headers=shop.h, json={
        "question": "What are your delivery areas?", "keywords": ["deliver", "Delivery"],
        "answer": "We deliver across Coimbatore and Tiruppur, Monday to Saturday."})
    assert r.status_code == 201 and r.json()["keywords"] == ["deliver", "delivery"]
    sim_text(shop.client, shop.h, NUM, "do you deliver to my area?")
    assert last_out(db, shop.c["id"]).text.startswith("We deliver across Coimbatore")
    assert shop.client.post(f"{API}/faqs/test", headers=shop.h, json={"text": "delivery areas?"}).json()["faq_id"]
    # not in the FAQ: handed to staff, flagged in the inbox, bot steps back
    sim_text(shop.client, shop.h, NUM, "do you sell cement bags?")
    assert "passed this to our team" in last_out(db, shop.c["id"]).text
    assert shop.client.get(f"{API}/inbox/attention", headers=shop.h).json() == {"count": 1}
    conv = shop.client.get(f"{API}/inbox/conversations", headers=shop.h, params={"attention": True}).json()["items"][0]
    assert conv["handover_reason"] == "Question the bot couldn't answer" and not conv["bot_enabled"]
    # staff reply clears the flag
    r = shop.client.post(f"{API}/inbox/conversations/{conv['id']}/messages", headers=shop.h,
                         json={"kind": "text", "text": "No sir, only FMCG."})
    assert r.status_code == 201
    assert shop.client.get(f"{API}/inbox/attention", headers=shop.h).json() == {"count": 0}
    # "ok thanks" is not a question: no handover
    shop.client.patch(f"{API}/inbox/conversations/{conv['id']}", headers=shop.h, json={"bot_enabled": True})
    sim_text(shop.client, shop.h, NUM, "ok thanks")
    assert shop.client.get(f"{API}/inbox/attention", headers=shop.h).json() == {"count": 0}


def test_buyer_asks_for_a_person(shop, db):  # noqa: F811
    sim_text(shop.client, shop.h, NUM, "I want to talk to a person")
    assert "passed this to our team" in last_out(db, shop.c["id"]).text
    before = last_out(db, shop.c["id"]).id
    sim_text(shop.client, shop.h, NUM, "hello?")  # bot is off for this chat now
    assert last_out(db, shop.c["id"]).id == before
    conv = db.scalar(select(Conversation).where(Conversation.customer_id == shop.c["id"]))
    r = shop.client.patch(f"{API}/inbox/conversations/{conv.id}", headers=shop.h, json={"bot_enabled": False, "resolved": True})
    assert r.json()["handover_at"] is None and r.json()["bot_enabled"] is False


# ================================================================== business assistant data
@pytest.fixture
def books(client, company_a, db):
    """Known data set. Placed today: Anand 2 x 1770, Bala 885, Elango 177 (on hold) = 4 orders, 4602.
    Outstanding: Anand 3540 (1770 is 70 days overdue, 1770 due in 3 days), Bala 585 (10 days overdue, paid 300),
    Chitra 1000 opening balance, no orders. Sales user Ravi has only Anand."""
    h = company_a.h(Role.owner)
    client.patch(f"{API}/company", headers=h, json={"gstin": "33AAACP1234F1Z5", "state": "Tamil Nadu"})
    connect_simulator(client, h)
    oil = product(client, h, "OIL-1L", 150, gst=18, name="Sunflower Oil 1L")
    product(client, h, "SOAP", 30, name="Bath Soap", stock_status="out_of_stock")
    sales = company_a.users[Role.sales]
    a = customer(client, h, number="9800000001", business_name="Anand Stores", whatsapp_opt_in=True,
                 assigned_user_id=str(sales.id), credit_days=30)
    b = customer(client, h, number="9800000002", business_name="Bala Traders", whatsapp_opt_in=True, credit_days=30)
    c = customer(client, h, number="9800000003", business_name="Chitra Agencies", opening_balance=1000)
    e = customer(client, h, number="9800000005", business_name="Elango Mart", credit_limit=100)
    d = customer(client, h, number="9800000004", business_name="Dormant Dealers")
    assert client.patch(f"{API}/customers/{d['id']}", headers=h, json={"status": "inactive"}).status_code == 200

    def order(cust, qty, approve=True):
        r = client.post(f"{API}/orders", headers=h, json={"customer_id": cust["id"], "action": "confirm",
                                                          "lines": [{"product_id": oil["id"], "quantity": qty}]})
        assert r.status_code == 201, r.text
        o = r.json()
        if approve and o["status"] == "confirmed":
            o = client.post(f"{API}/orders/{o['id']}/approve", headers=h, json={}).json()
        return o

    s = SimpleNamespace(client=client, h=h, company=company_a, a=a, b=b, c=c, e=e)
    inv = {}
    for key, cust, qty, due in (("a1", a, 10, -70), ("a2", a, 10, 3), ("b1", b, 5, -10)):
        o = make_invoice(s, order(cust, qty))
        inv[key] = o["invoice_id"]
        db.execute(update(Invoice).where(Invoice.id == o["invoice_id"]).values(due_date=an.today() + timedelta(days=due)))
    db.commit()
    assert order(e, 1, approve=False)["status"] == "on_hold"
    r = client.post(f"{API}/payments", headers=h, json={"customer_id": b["id"], "amount": 300, "mode": "cash"})
    assert r.status_code == 201, r.text
    sim_text(client, h, "9800000001", "hi")  # one chat waiting for staff
    s.inv = inv
    return s


def ask(b, q, headers=None):
    r = b.client.post(f"{API}/assistant/ask", headers=headers or b.h, json={"question": q})
    assert r.status_code == 200, r.text
    return r.json()


QUESTIONS = [
    ("who is above 60 days overdue?", "1 customer has invoices more than 60 days overdue, ₹1,770.00 in total"),
    ("which customers are overdue", "2 customers have invoices overdue, ₹2,355.00 in total"),
    ("top 10 customers this month", "Anand Stores (₹3,540.00)"),
    ("customers with no order in 30 days", "1 customer has not ordered in the last 30 days (1 never ordered)."),
    ("best selling products this month", "Sunflower Oil 1L (₹4,602.00)"),
    ("how many orders this month?", "4 orders worth ₹4,602.00 this month"),
    ("sales today", "4 orders worth ₹4,602.00 today"),
    ("orders yesterday", "No orders were placed yesterday."),
    ("how much did we collect this month", "₹300.00 collected this month from 1 payment."),
    ("total outstanding", "Total outstanding is ₹5,125.00 across 3 customers. Highest: Anand Stores (₹3,540.00)."),
    ("how much does Bala Traders owe?", "Bala Traders owes ₹585.00, of which ₹585.00 is overdue"),
    ("Anand Stores' outstanding", "Anand Stores owes ₹3,540.00, of which ₹1,770.00 is overdue"),
    ("orders pending approval", "1 order worth ₹177.00 is waiting for approval"),
    ("invoices due this week", "1 invoice worth ₹1,770.00 falls due in the next 7 days."),
    ("what is out of stock?", "1 product out of stock and 0 running low."),
    ("unanswered chats", "1 chat waiting for a reply"),
    ("orders of Bala Traders this month", "Bala Traders: 1 order placed this month, worth ₹885.00"),
    ("who has not ordered in the last 60 days", "1 customer has not ordered in the last 60 days"),
    ("top 2 customers by order value last 7 days", "Top 2 customers by order value in the last 7 days"),
    ("payments received today", "₹300.00 collected today"),
]


def test_twenty_questions_with_known_answers(books):
    """AC: at least 18 of 20 correct. Every figure comes from the database, never from the AI."""
    right = []
    for q, expected in QUESTIONS:
        a = ask(books, q)
        ok = a["kind"] == "answer" and expected in a["answer"]
        right.append(ok)
        if not ok:
            print("MISS", q, "->", a["tool"], a["answer"])
    assert sum(right) >= 18
    assert all(right)  # the rule-based router gets all 20; a real LLM must still get 18


def test_unknown_question_is_not_guessed(books):
    a = ask(books, "what's the weather in Chennai?")
    assert a["kind"] == "unknown" and "won't guess" in a["answer"]


def test_answers_respect_role(books):
    """FR-M10-07: Sales only gets their own customers."""
    sales_h = books.company.h(Role.sales)
    assert "Total outstanding is ₹3,540.00 across 1 customer" in ask(books, "total outstanding", sales_h)["answer"]
    assert "couldn't find a customer called \"Bala Traders\" among your customers" in ask(
        books, "how much does Bala Traders owe?", sales_h)["answer"]
    a = ask(books, "top 10 customers this month", sales_h)
    assert [r["customer"] for r in a["rows"]] == ["Anand Stores"]
    d = books.client.get(f"{API}/dashboard", headers=sales_h).json()
    assert d["orders"]["count"] == 2 and d["scope"] == "own customers"


def test_table_export_to_excel(books):
    a = ask(books, "which customers are overdue")
    assert [c["label"] for c in a["columns"]][:2] == ["Customer", "Overdue"] and len(a["rows"]) == 2
    r = books.client.get(f"{API}/assistant/messages/{a['id']}/export", headers=books.h)
    assert r.status_code == 200
    ws = load_workbook(BytesIO(r.content)).active
    assert [c.value for c in ws[1]][:2] == ["Customer", "Overdue"]
    assert ws["A2"].value == "Anand Stores" and ws["B2"].value == 1770.0
    hist = books.client.get(f"{API}/assistant/messages", headers=books.h).json()
    assert hist[-1]["id"] == a["id"]
    # other users can't read or export it
    assert books.client.get(f"{API}/assistant/messages/{a['id']}/export",
                            headers=books.company.h(Role.admin)).status_code == 404


def test_owner_rule_needs_confirmation_then_runs_mondays(books, db):
    """FR-M10-09."""
    admin = ask(books, "remind everyone above 60 days every Monday", books.company.h(Role.admin))
    assert admin["kind"] == "answer" and admin["answer"].startswith("Only the Owner")
    p = ask(books, "remind everyone above 60 days every Monday at 10am")
    assert p["kind"] == "rule_proposal"
    assert "more than 60 days overdue, every Monday at 10 AM IST" in p["answer"]
    assert db.scalars(select(AutomationRule)).all() == []  # nothing until confirmed
    r = books.client.post(f"{API}/assistant/messages/{p['id']}/confirm-rule", headers=books.h)
    assert r.status_code == 200, r.text
    assert books.client.post(f"{API}/assistant/messages/{p['id']}/confirm-rule", headers=books.h).status_code == 409
    rule = r.json()
    assert rule["weekdays"] == [0] and rule["hour"] == 10 and rule["params"] == {"min_days": 60}
    # Tuesday: nothing; Monday 09:30: not yet; Monday 10:05: Anand (70 days) is reminded, Bala (10 days) is not
    tue = datetime(2026, 10, 27, 10, 5, tzinfo=an.IST)
    mon_early = datetime(2026, 10, 26, 9, 30, tzinfo=an.IST)
    mon = datetime(2026, 10, 26, 10, 5, tzinfo=an.IST)
    assert automation.run_due_rules(db, tue) == 0 and automation.run_due_rules(db, mon_early) == 0
    assert automation.run_due_rules(db, mon) == 1
    order_svc.commit_and_send(db)
    assert automation.run_due_rules(db, mon + timedelta(hours=2)) == 0  # once a day
    a_msg = last_out(db, books.a["id"]).text
    assert "payment reminder" in a_msg and "Total due: ₹3,540.00" in a_msg
    b_last = last_out(db, books.b["id"])
    assert b_last is None or "payment reminder" not in (b_last.text or "")
    rules = books.client.get(f"{API}/assistant/rules", headers=books.h).json()
    assert rules[0]["last_run_summary"].endswith("1 reminder(s) sent, 0 skipped")
    assert books.client.patch(f"{API}/assistant/rules/{rule['id']}", headers=books.h,
                              json={"is_active": False}).json()["is_active"] is False


def test_daily_summary_at_set_time_once_a_day(books, db):
    """FR-M10-08 / AC: arrives within 5 minutes of the set time (beat runs every minute), once per day."""
    tid = books.company.tenant_id
    books.client.patch(f"{API}/company", headers=books.h, json={"daily_summary_time": "09:00"})
    tomorrow = an.today() + timedelta(days=1)
    before = datetime.combine(tomorrow, datetime.min.time(), an.IST).replace(hour=8, minute=59)
    at = before.replace(hour=9, minute=1)
    assert automation.run_due_summaries(db, before) == 0
    assert automation.run_due_summaries(db, at) == 1
    assert automation.run_due_summaries(db, at + timedelta(minutes=5)) == 0
    s = db.scalar(select(DailySummary).where(DailySummary.tenant_id == tid))
    owner = db.get(User, books.company.users[Role.owner].id)
    assert s.sent_to == owner.mobile and s.sent_at is not None and s.send_error is None
    assert s.data["orders"] == 4 and s.data["order_value"] == "4602.00" and s.data["collections"] == "300.00"
    assert s.data["pending_approvals"] == 1 and s.data["unanswered_chats"] == 1
    assert "• Orders: 4 worth ₹4,602.00" in s.text and "• Collections: ₹300.00" in s.text
    # shown in the panel; can be re-sent now
    assert books.client.get(f"{API}/assistant/summaries", headers=books.h).json()[0]["text"] == s.text
    assert books.client.post(f"{API}/assistant/summaries/send-now", headers=books.h).status_code == 200
    # disabled -> not sent
    db.execute(update(Tenant).where(Tenant.id == tid).values(daily_summary_enabled=False))
    db.commit()
    assert automation.run_due_summaries(db, at + timedelta(days=1)) == 0


def test_dashboard_figures(books):
    d = books.client.get(f"{API}/dashboard", headers=books.h, params={"period": "this_month"}).json()
    assert d["orders"]["count"] == 4 and d["orders"]["value"] == "4602.00"
    by = {r["source"]: r["orders"] for r in d["orders"]["by_source"]}
    assert by == {"WhatsApp (AI)": 0, "Panel": 4, "Excel": 0}
    assert d["collections"]["amount"] == "300.00"
    assert d["outstanding"]["total"] == "5125.00" and d["outstanding"]["overdue"] == "2355.00"
    assert d["top_customers"][0]["customer"] == "Anand Stores"
    assert d["top_products"][0]["product"] == "Sunflower Oil 1L"
    assert d["series"][-1]["date"] == an.today().isoformat() and d["series"][-1]["orders"] == 4
    bad = books.client.get(f"{API}/dashboard", headers=books.h, params={"start": "2026-01-10", "end": "2026-01-01"})
    assert bad.status_code == 400


def test_period_parsing():
    on = an.period("today").start
    assert an.period_from_text("sales last 15 days").label == "in the last 15 days"
    assert an.period_from_text("orders in march").start.month == 3
    assert an.period("last_week").end.weekday() == 6
    assert an.inr(1234567.5) == "₹12,34,567.50" and an.inr(-500) == "-₹500.00"
    assert an.period("this_month", on=on).start.day == 1
