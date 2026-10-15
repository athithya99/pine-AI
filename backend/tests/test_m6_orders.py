"""M6 acceptance criteria and order rules."""

from datetime import UTC, date, datetime, timedelta

from sqlalchemy import func, select

from app.models import Direction, Message, Order, Role
from app.services import orders as svc
from tests.test_m2_customers import upload, xlsx
from tests.test_m3_catalog import customer, product
from tests.test_m4_whatsapp import connect_simulator, sim_text

API = "/api/v1"
NUM = "9876500001"


def make_order(client, h, customer_id, lines, action="draft", expect=201):
    r = client.post(f"{API}/orders", headers=h, json={
        "customer_id": customer_id, "action": action,
        "lines": [{"product_id": p, "quantity": q} for p, q in lines]})
    assert r.status_code == expect, r.text
    return r.json()


def act(client, h, order_id, action, body=None, expect=200):
    r = client.post(f"{API}/orders/{order_id}/{action}", headers=h, json=body or {})
    assert r.status_code == expect, r.text
    return r.json()


def outbound_count(db, customer_id):
    db.expire_all()
    return db.scalar(select(func.count()).select_from(Message).where(
        Message.customer_id == customer_id, Message.direction == Direction.outbound))


def last_outbound(db, customer_id):
    db.expire_all()
    return db.scalars(select(Message).where(Message.customer_id == customer_id, Message.direction == Direction.outbound)
                      .order_by(Message.created_at.desc())).first()


def test_order_number_format_and_fy_reset():
    assert svc.financial_year(date(2026, 10, 15)) == "26-27"
    assert svc.financial_year(date(2027, 3, 31)) == "26-27"
    assert svc.financial_year(date(2027, 4, 1)) == "27-28"
    assert svc.financial_year(date(2026, 5, 1), start_month=1) == "26"


def test_create_order_prices_from_m3_and_unique_numbers(client, company_a, company_b):
    """FR-M6-02/04: prices follow M3; numbers are per tenant."""
    h = company_a.h(Role.owner)
    p = product(client, h, "SALT", 50, gst=5)
    c = customer(client, h)
    client.put(f"{API}/customers/{c['id']}/prices", headers=h, json={"items": [{"product_id": p["id"], "price": 45}]})
    o1 = make_order(client, h, c["id"], [(p["id"], 10)])
    o2 = make_order(client, h, c["id"], [(p["id"], 1)])
    assert o1["status"] == "draft" and o1["number"].startswith("ORD/") and o1["number"].endswith("/00001")
    assert o2["number"].endswith("/00002")
    line = o1["lines"][0]
    assert line["unit_price"] == "45.00" and line["price_source"] == "customer_price"
    assert o1["taxable_total"] == "450.00" and o1["gst_total"] == "22.50" and o1["grand_total"] == "472.50"

    hb = company_b.h(Role.owner)
    pb = product(client, hb, "SALT", 50)
    cb = customer(client, hb)
    ob = make_order(client, hb, cb["id"], [(pb["id"], 1)])
    assert ob["number"].endswith("/00001")
    # tenant isolation
    assert client.get(f"{API}/orders/{o1['id']}", headers=hb).status_code == 404


def test_custom_prefix(client, company_a):
    h = company_a.h(Role.owner)
    r = client.patch(f"{API}/company", headers=h, json={"order_prefix": "pso", "quote_validity_hours": 12})
    assert r.status_code == 200 and r.json()["order_prefix"] == "PSO"
    p = product(client, h, "A", 10)
    c = customer(client, h)
    assert make_order(client, h, c["id"], [(p["id"], 1)])["number"].startswith("PSO/")


def test_credit_limit_puts_order_on_hold(client, company_a):
    """AC: ₹50,000 limit, ₹45,000 outstanding, ₹10,000 order -> On hold, in the owner's approval queue."""
    h = company_a.h(Role.owner)
    p = product(client, h, "BIG", 10000, gst=0)
    c = customer(client, h, credit_limit=50000, opening_balance=45000)
    o = make_order(client, h, c["id"], [(p["id"], 1)], action="confirm")
    assert o["status"] == "on_hold"
    assert "credit limit" in o["hold_reasons"].lower()
    assert o["credit"]["outstanding"] == "45000.00"

    queue = client.get(f"{API}/orders", headers=h, params={"status": "on_hold"}).json()
    assert [x["id"] for x in queue["items"]] == [o["id"]]
    assert client.get(f"{API}/orders/counts", headers=h).json()["on_hold"] == 1

    # Sales cannot override the credit limit; approval with a remark by the owner works
    sales = company_a.h(Role.sales)
    assert client.post(f"{API}/orders/{o['id']}/approve", headers=sales, json={"remark": "ok"}).status_code == 403
    act(client, h, o["id"], "approve", {}, expect=400)  # remark required for on-hold
    o = act(client, h, o["id"], "approve", {"remark": "Payment promised Friday"})
    assert o["status"] == "approved" and o["approval_remark"] == "Payment promised Friday"


def test_within_credit_limit_is_confirmed(client, company_a):
    h = company_a.h(Role.owner)
    p = product(client, h, "SMALL", 1000, gst=0)
    c = customer(client, h, credit_limit=50000, opening_balance=45000)
    o = make_order(client, h, c["id"], [(p["id"], 4)], action="confirm")
    assert o["status"] == "confirmed"
    # approved orders count towards what the buyer owes
    act(client, h, o["id"], "approve")
    o2 = make_order(client, h, c["id"], [(p["id"], 2)], action="confirm")
    assert o2["status"] == "on_hold"  # 45,000 + 4,000 + 2,000 > 50,000


def test_reject_on_hold(client, company_a):
    h = company_a.h(Role.owner)
    p = product(client, h, "OOS", 100)
    client.put(f"{API}/products/{p['id']}/stock", headers=h, json={"stock_status": "out_of_stock"})
    c = customer(client, h)
    o = make_order(client, h, c["id"], [(p["id"], 1)], action="confirm")
    assert o["status"] == "on_hold" and "Out of stock" in o["hold_reasons"]
    act(client, h, o["id"], "reject", {"remark": ""}, expect=400)
    o = act(client, h, o["id"], "reject", {"remark": "No stock till next month"})
    assert o["status"] == "cancelled" and "No stock" in o["cancel_reason"]


def test_invalid_transitions_rejected(client, company_a):
    """AC: Quoted -> Dispatched is rejected."""
    h = company_a.h(Role.owner)
    p = product(client, h, "A", 10)
    c = customer(client, h)
    o = make_order(client, h, c["id"], [(p["id"], 1)], action="quote")
    assert o["status"] == "quoted"
    r = client.post(f"{API}/orders/{o['id']}/dispatch", headers=h, json={"vehicle_number": "TN37AB1234"})
    assert r.status_code == 409 and "Quoted to Dispatched" in r.json()["detail"]
    assert client.post(f"{API}/orders/{o['id']}/deliver", headers=h).status_code == 409
    assert client.post(f"{API}/orders/{o['id']}/approve", headers=h, json={}).status_code == 409
    assert client.post(f"{API}/orders/{o['id']}/invoice", headers=h, json={"invoice_number": "1"}).status_code == 409
    assert client.get(f"{API}/orders/{o['id']}", headers=h).json()["status"] == "quoted"


def test_each_buyer_facing_change_sends_exactly_one_message(client, company_a, db):
    """AC: each buyer-facing status change sends exactly one WhatsApp message; Draft and On hold send none."""
    h = company_a.h(Role.owner)
    connect_simulator(client, h)
    p = product(client, h, "RICE", 60, gst=5)
    c = customer(client, h)
    sim_text(client, h, NUM, "hi")  # opens the 24-hour window
    base = outbound_count(db, c["id"])

    o = make_order(client, h, c["id"], [(p["id"], 5)])
    assert outbound_count(db, c["id"]) == base  # draft: no message

    o = act(client, h, o["id"], "send-quote")
    assert outbound_count(db, c["id"]) == base + 1
    quote_msg = last_outbound(db, c["id"])
    buttons = quote_msg.payload["interactive"]["action"]["buttons"]
    assert [b["reply"]["title"] for b in buttons] == ["Confirm", "Edit", "Cancel"]
    assert o["number"] in quote_msg.text and "₹315.00" in quote_msg.text

    # buyer taps Confirm
    client.post(f"{API}/simulator/inbound", headers=h,
                json={"from_number": NUM, "reply_id": f"order:confirm:{o['id']}", "reply_title": "Confirm"})
    o = client.get(f"{API}/orders/{o['id']}", headers=h).json()
    assert o["status"] == "confirmed" and o["confirmed_by_buyer"] is True
    assert outbound_count(db, c["id"]) == base + 2

    o = act(client, h, o["id"], "approve")
    assert outbound_count(db, c["id"]) == base + 3
    o = act(client, h, o["id"], "invoice", {"invoice_number": "INV-101"})
    assert outbound_count(db, c["id"]) == base + 4
    o = act(client, h, o["id"], "dispatch", {"vehicle_number": "tn37ab1234", "lr_number": "LR55"})
    assert o["vehicle_number"] == "TN37AB1234"
    assert outbound_count(db, c["id"]) == base + 5
    assert "TN37AB1234" in last_outbound(db, c["id"]).text
    o = act(client, h, o["id"], "deliver")
    assert o["status"] == "delivered"
    assert outbound_count(db, c["id"]) == base + 6
    sent = [e for e in o["events"] if e["kind"] == "status"]
    assert all(e["message_id"] for e in sent)

    # a second tap on Confirm is answered, not re-processed
    client.post(f"{API}/simulator/inbound", headers=h,
                json={"from_number": NUM, "reply_id": f"order:confirm:{o['id']}", "reply_title": "Confirm"})
    assert "already delivered" in last_outbound(db, c["id"]).text


def test_on_hold_sends_only_the_confirmation(client, company_a, db):
    h = company_a.h(Role.owner)
    connect_simulator(client, h)
    p = product(client, h, "BIG", 10000, gst=0)
    c = customer(client, h, credit_limit=5000)
    sim_text(client, h, NUM, "hi")
    base = outbound_count(db, c["id"])
    o = make_order(client, h, c["id"], [(p["id"], 1)], action="confirm")
    assert o["status"] == "on_hold"
    assert outbound_count(db, c["id"]) == base + 1
    assert "received your order" in last_outbound(db, c["id"]).text


def test_buyer_cancel_and_edit_buttons(client, company_a, db):
    h = company_a.h(Role.owner)
    connect_simulator(client, h)
    p = product(client, h, "A", 10)
    c = customer(client, h)
    sim_text(client, h, NUM, "hi")
    o = make_order(client, h, c["id"], [(p["id"], 1)], action="quote")
    client.post(f"{API}/simulator/inbound", headers=h, json={"from_number": NUM, "reply_id": f"order:edit:{o['id']}"})
    assert "what you'd like to change" in last_outbound(db, c["id"]).text
    assert client.get(f"{API}/orders/{o['id']}", headers=h).json()["status"] == "quoted"
    client.post(f"{API}/simulator/inbound", headers=h, json={"from_number": NUM, "reply_id": f"order:cancel:{o['id']}"})
    o = client.get(f"{API}/orders/{o['id']}", headers=h).json()
    assert o["status"] == "cancelled" and o["cancelled_by_buyer"] is True
    assert "cancelled as requested" in last_outbound(db, c["id"]).text


def test_buyer_cannot_act_on_another_buyers_order(client, company_a, db):
    h = company_a.h(Role.owner)
    connect_simulator(client, h)
    p = product(client, h, "A", 10)
    c = customer(client, h)
    o = make_order(client, h, c["id"], [(p["id"], 1)], action="quote")
    other = "9876500002"
    sim_text(client, h, other, "hi")
    client.post(f"{API}/simulator/inbound", headers=h, json={"from_number": other, "reply_id": f"order:confirm:{o['id']}"})
    assert client.get(f"{API}/orders/{o['id']}", headers=h).json()["status"] == "quoted"


def test_edit_after_confirm_requotes_and_is_audited(client, company_a, db):
    """FR-M6-08."""
    h = company_a.h(Role.owner)
    connect_simulator(client, h)
    p = product(client, h, "A", 10, gst=0)
    c = customer(client, h)
    sim_text(client, h, NUM, "hi")
    o = make_order(client, h, c["id"], [(p["id"], 1)], action="confirm")
    assert o["status"] == "confirmed"
    r = client.put(f"{API}/orders/{o['id']}/lines", headers=h, json={"lines": [{"product_id": p["id"], "quantity": 3}]})
    assert r.status_code == 200, r.text
    o = r.json()
    assert o["status"] == "quoted" and o["revision"] == 2 and o["grand_total"] == "30.00"
    assert last_outbound(db, c["id"]).text.startswith("Revised quote")
    logs = client.get(f"{API}/audit-logs", headers=h, params={"entity": "order", "entity_id": o["id"]}).json()
    assert any(l["action"] == "update" and "lines" in (l["new_values"] or {}) for l in logs)
    # can't edit after approval
    o = act(client, h, o["id"], "confirm")
    o = act(client, h, o["id"], "approve")
    r = client.put(f"{API}/orders/{o['id']}/lines", headers=h, json={"lines": [{"product_id": p["id"], "quantity": 1}]})
    assert r.status_code == 409


def test_quote_expiry(client, company_a, db):
    h = company_a.h(Role.owner)
    connect_simulator(client, h)
    p = product(client, h, "A", 10)
    c = customer(client, h)
    sim_text(client, h, NUM, "hi")
    o1 = make_order(client, h, c["id"], [(p["id"], 1)], action="quote")
    o2 = make_order(client, h, c["id"], [(p["id"], 2)], action="quote")
    past = datetime.now(UTC) - timedelta(minutes=1)
    for oid in (o1["id"], o2["id"]):
        db.get(Order, oid).quote_valid_until = past
    db.commit()

    # buyer taps Confirm on an expired quote -> Expired + told so
    client.post(f"{API}/simulator/inbound", headers=h, json={"from_number": NUM, "reply_id": f"order:confirm:{o1['id']}"})
    assert client.get(f"{API}/orders/{o1['id']}", headers=h).json()["status"] == "expired"
    assert "expired" in last_outbound(db, c["id"]).text

    # periodic job expires the rest
    from app.core.database import SessionLocal
    with SessionLocal() as s:
        assert svc.expire_due_quotes(s) == 1
        svc.commit_and_send(s)
    assert client.get(f"{API}/orders/{o2['id']}", headers=h).json()["status"] == "expired"
    # staff can re-quote an expired order
    o2 = act(client, h, o2["id"], "send-quote")
    assert o2["status"] == "quoted"


def test_where_is_my_order(client, company_a, db):
    """FR-M6-10."""
    h = company_a.h(Role.owner)
    connect_simulator(client, h)
    p = product(client, h, "A", 10)
    c = customer(client, h)
    o = make_order(client, h, c["id"], [(p["id"], 1)], action="confirm")
    o = act(client, h, o["id"], "approve")
    o = act(client, h, o["id"], "invoice", {"invoice_number": "INV-9"})
    act(client, h, o["id"], "dispatch", {"vehicle_number": "TN38X1", "lr_number": "77"})
    sim_text(client, h, NUM, "Where is my order?")
    reply = last_outbound(db, c["id"])
    assert reply.sender.value == "bot"
    assert o["number"] in reply.text and "Dispatched" in reply.text and "TN38X1" in reply.text

    # bot switched off (human takeover) -> no automatic reply
    conv = client.get(f"{API}/inbox/conversations", headers=h).json()["items"][0]
    client.patch(f"{API}/inbox/conversations/{conv['id']}", headers=h, json={"bot_enabled": False})
    before = outbound_count(db, c["id"])
    sim_text(client, h, NUM, "order status")
    assert outbound_count(db, c["id"]) == before


def test_status_query_matching():
    for t in ["where is my order", "Order status?", "status", "my orders", "order enga iruku", "mera order kahan hai"]:
        assert svc.is_status_query(t), t
    for t in ["10 bags salt", "hi", "price of rice"]:
        assert not svc.is_status_query(t), t


def test_no_window_no_template_still_moves_and_records_why(client, company_a):
    h = company_a.h(Role.owner)
    connect_simulator(client, h)
    p = product(client, h, "A", 10)
    c = customer(client, h)  # never messaged us, no template
    o = make_order(client, h, c["id"], [(p["id"], 1)], action="quote")
    assert o["status"] == "quoted"
    ev = [e for e in o["events"] if e["to_status"] == "quoted"][0]
    assert ev["message_id"] is None and "order_update" in ev["notify_note"]


def test_template_used_outside_window(client, company_a, db):
    from app.models import MessageTemplate, TemplateCategory, TemplateStatus

    h = company_a.h(Role.owner)
    connect_simulator(client, h)
    db.add(MessageTemplate(tenant_id=company_a.tenant_id, name="order_update", language="en",
                           category=TemplateCategory.utility, body="Hi {{1}}, update on order {{2}}.",
                           variable_count=2, status=TemplateStatus.approved))
    db.commit()
    p = product(client, h, "A", 10)
    c = customer(client, h, contact_person="Murugan", whatsapp_opt_in=True)
    o = make_order(client, h, c["id"], [(p["id"], 1)], action="confirm")
    msg = last_outbound(db, c["id"])
    assert msg.msg_type.value == "template" and o["number"] in msg.text and "Murugan" in msg.text


def test_short_supply_on_invoice(client, company_a):
    """FR-M6-09."""
    h = company_a.h(Role.owner)
    p1 = product(client, h, "A", 10)
    p2 = product(client, h, "B", 20)
    c = customer(client, h)
    o = make_order(client, h, c["id"], [(p1["id"], 10), (p2["id"], 5)], action="confirm")
    o = act(client, h, o["id"], "approve")
    line_a = next(l for l in o["lines"] if l["sku"] == "A")
    act(client, h, o["id"], "invoice", {"invoice_number": "X", "lines": [{"line_id": line_a["id"], "invoiced_qty": 11}]},
        expect=400)
    o = act(client, h, o["id"], "invoice", {"invoice_number": "INV-1",
                                             "lines": [{"line_id": line_a["id"], "invoiced_qty": 8}]})
    lines = {l["sku"]: l for l in o["lines"]}
    assert o["status"] == "invoiced" and lines["A"]["invoiced_qty"] == "8.000" and lines["B"]["invoiced_qty"] == "5.000"
    assert "short supply" in [e for e in o["events"] if e["to_status"] == "invoiced"][0]["remark"]


def test_role_permissions(client, company_a):
    h = company_a.h(Role.owner)
    p = product(client, h, "A", 10)
    sales_user = company_a.users[Role.sales]
    mine = customer(client, h, assigned_user_id=str(sales_user.id))
    theirs = customer(client, h, number="9876500009")
    sales, accounts = company_a.h(Role.sales), company_a.h(Role.accounts)

    o_mine = make_order(client, sales, mine["id"], [(p["id"], 1)])
    assert o_mine["sales_user_id"] == str(sales_user.id)
    make_order(client, sales, theirs["id"], [(p["id"], 1)], expect=404)
    o_theirs = make_order(client, h, theirs["id"], [(p["id"], 1)])
    assert client.get(f"{API}/orders/{o_theirs['id']}", headers=sales).status_code == 404
    assert [x["id"] for x in client.get(f"{API}/orders", headers=sales).json()["items"]] == [o_mine["id"]]

    # accounts: view everything, can't create, can invoice
    assert client.get(f"{API}/orders", headers=accounts).json()["total"] == 2
    make_order(client, accounts, mine["id"], [(p["id"], 1)], expect=403)
    o = act(client, sales, o_mine["id"], "confirm")
    assert "approve" not in o["allowed_actions"]
    act(client, h, o_mine["id"], "approve")
    assert client.post(f"{API}/orders/{o_mine['id']}/cancel", headers=sales, json={"reason": "x y"}).status_code == 403
    o = act(client, accounts, o_mine["id"], "invoice", {"invoice_number": "INV-2"})
    assert o["status"] == "invoiced"
    assert client.post(f"{API}/orders/{o_mine['id']}/dispatch", headers=accounts, json={}).status_code == 403


def test_cancel_needs_reason_and_inactive_product_rejected(client, company_a):
    h = company_a.h(Role.owner)
    p = product(client, h, "A", 10)
    c = customer(client, h)
    o = make_order(client, h, c["id"], [(p["id"], 1)])
    assert client.post(f"{API}/orders/{o['id']}/cancel", headers=h, json={"reason": ""}).status_code == 422
    o = act(client, h, o["id"], "cancel", {"reason": "Duplicate order"})
    assert o["status"] == "cancelled" and o["allowed_actions"] == []
    client.patch(f"{API}/products/{p['id']}", headers=h, json={"is_active": False})
    r = client.post(f"{API}/orders", headers=h, json={"customer_id": c["id"], "lines": [{"product_id": p["id"], "quantity": 1}]})
    assert r.status_code == 400 and "Inactive" in r.json()["detail"]


def test_filters_search_and_export(client, company_a):
    h = company_a.h(Role.owner)
    p = product(client, h, "A", 10)
    c1 = customer(client, h, business_name="Murugan Stores")
    c2 = customer(client, h, number="9876500002", business_name="Lakshmi Traders")
    o1 = make_order(client, h, c1["id"], [(p["id"], 1)], action="quote")
    make_order(client, h, c2["id"], [(p["id"], 2)])
    assert client.get(f"{API}/orders", headers=h, params={"q": "lakshmi"}).json()["total"] == 1
    assert client.get(f"{API}/orders", headers=h, params={"q": o1["number"]}).json()["items"][0]["id"] == o1["id"]
    assert client.get(f"{API}/orders", headers=h, params={"status": ["quoted", "draft"]}).json()["total"] == 2
    assert client.get(f"{API}/orders", headers=h, params={"customer_id": c1["id"]}).json()["total"] == 1
    assert client.get(f"{API}/orders", headers=h, params={"source": "excel"}).json()["total"] == 0
    today = datetime.now(svc.IST).date().isoformat()
    assert client.get(f"{API}/orders", headers=h, params={"date_from": today, "date_to": today}).json()["total"] == 2
    r = client.get(f"{API}/orders/export", headers=h)
    assert r.status_code == 200
    from io import BytesIO

    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["Orders", "Order lines"]
    assert wb["Orders"].max_row == 3 and wb["Order lines"].max_row == 3


def test_excel_import(client, company_a):
    """FR-M6-03: valid orders become drafts; an order with any bad row is skipped whole."""
    h = company_a.h(Role.owner)
    product(client, h, "SALT", 20)
    product(client, h, "RICE", 60)
    customer(client, h)
    content = xlsx(["order_ref", "customer_whatsapp", "sku", "quantity", "notes"], [
        ["A1", "9876500001", "SALT", 10, "Urgent"],
        ["A1", "9876500001", "rice", 2, None],
        ["B1", "9876500001", "SALT", 5, None],
        ["B1", "9876500001", "NOPE", 1, None],
        ["C1", "9999999999", "SALT", 1, None],
    ])
    r = upload(client, h, f"{API}/orders/import", content)
    assert r.status_code == 200, r.text
    res = r.json()
    assert res["created"] == 1 and res["total_rows"] == 5
    assert {e["row"] for e in res["errors"]} == {5, 6}
    orders = client.get(f"{API}/orders", headers=h).json()["items"]
    assert len(orders) == 1 and orders[0]["source"] == "excel" and orders[0]["line_count"] == 2
    detail = client.get(f"{API}/orders/{orders[0]['id']}", headers=h).json()
    assert detail["notes"] == "Urgent" and detail["grand_total"] == "336.00"
    assert client.get(f"{API}/orders/import/template", headers=h).status_code == 200
