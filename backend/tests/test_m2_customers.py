"""M2 acceptance criteria."""

from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.models import Role

API = "/api/v1"


def xlsx(headers, rows) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(r)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def upload(client, headers, path, content):
    return client.post(path, headers=headers, files={"file": ("data.xlsx", content, "application/octet-stream")})


def test_import_1000_rows_with_10_invalid(client, company_a):
    """AC: 1,000 rows with 10 invalid -> 990 created, report names the 10 rows."""
    rows = []
    bad_rows = set()
    for i in range(1000):
        number = f"98765{i:05d}"
        name = f"Shop {i}"
        if i % 100 == 7:          # 10 bad rows: alternate between bad phone and missing name
            if (i // 100) % 2:
                number = "12345"
            else:
                name = None
            bad_rows.add(i + 2)    # +1 header, +1 one-based
        rows.append([name, number, "retailer", "coimbatore", "yes", 25000, 30])
    content = xlsx(["business_name*", "whatsapp_number*", "customer_type", "tags", "whatsapp_opt_in",
                    "credit_limit", "credit_days"], rows)

    r = upload(client, company_a.h(Role.owner), f"{API}/customers/import", content)
    assert r.status_code == 200, r.text
    result = r.json()
    assert result["total_rows"] == 1000
    assert result["created"] == 990
    assert result["failed"] == 10
    assert {e["row"] for e in result["errors"]} == bad_rows
    assert all(e["errors"] for e in result["errors"])

    listing = client.get(f"{API}/customers?size=1", headers=company_a.h(Role.owner)).json()
    assert listing["total"] == 990
    one = client.get(f"{API}/customers?q=Shop 1&size=1", headers=company_a.h(Role.owner)).json()["items"][0]
    assert one["whatsapp_opt_in"] is True and one["opt_in_source"] == "excel_import"


def test_whatsapp_number_unique_per_tenant(client, company_a, company_b):
    """AC: two customers in the same tenant cannot share a WhatsApp number."""
    body = {"business_name": "First", "whatsapp_number": "9876500001"}
    assert client.post(f"{API}/customers", headers=company_a.h(Role.owner), json=body).status_code == 201
    # same number written differently is still a duplicate
    dup = {"business_name": "Second", "whatsapp_number": "+91 98765 00001"}
    assert client.post(f"{API}/customers", headers=company_a.h(Role.owner), json=dup).status_code == 409
    # another tenant may have the same buyer
    assert client.post(f"{API}/customers", headers=company_b.h(Role.owner), json=dup).status_code == 201


def test_import_reports_duplicates_and_unknown_refs(client, company_a):
    client.post(f"{API}/customers", headers=company_a.h(Role.owner),
                json={"business_name": "Existing", "whatsapp_number": "9876500001"})
    content = xlsx(
        ["business_name", "whatsapp_number", "price_list", "sales_user_email", "gstin"],
        [
            ["A Shop", "9876500001", None, None, None],          # exists already
            ["B Shop", "9876500002", None, None, None],          # ok
            ["C Shop", "9876500002", None, None, None],          # duplicate in file
            ["D Shop", "9876500003", "Nope", None, None],        # unknown price list
            ["E Shop", "9876500004", None, "ghost@x-demo.in", None],  # unknown sales user
            ["F Shop", "9876500005", None, None, "BADGSTIN"],    # bad GSTIN
        ],
    )
    result = upload(client, company_a.h(Role.owner), f"{API}/customers/import", content).json()
    assert result["created"] == 1
    assert [e["row"] for e in result["errors"]] == [2, 4, 5, 6, 7]


def test_import_rejects_wrong_file(client, company_a):
    r = upload(client, company_a.h(Role.owner), f"{API}/customers/import", b"not an excel file")
    assert r.status_code == 400
    r = upload(client, company_a.h(Role.owner), f"{API}/customers/import", xlsx(["name"], [["x"]]))
    assert r.status_code == 400 and "Missing required columns" in r.json()["detail"]


def test_template_and_export(client, company_a):
    owner = company_a.h(Role.owner)
    t = client.get(f"{API}/customers/import/template", headers=owner)
    assert t.status_code == 200
    wb = load_workbook(BytesIO(t.content))
    assert wb.active.cell(1, 1).value == "business_name*"

    client.post(f"{API}/customers", headers=owner,
                json={"business_name": "Export Me", "whatsapp_number": "9876500001", "tags": ["Coimbatore", " coimbatore"]})
    e = client.get(f"{API}/customers/export", headers=owner)
    ws = load_workbook(BytesIO(e.content)).active
    assert ws.cell(2, 1).value == "Export Me"
    assert ws.cell(2, 13).value == "coimbatore"  # tags normalised and de-duplicated

    # the exported file can be re-imported into another company
    # (reuses the same layout as the template)


def test_opt_in_tracking(client, company_a):
    owner = company_a.h(Role.owner)
    c = client.post(f"{API}/customers", headers=owner,
                    json={"business_name": "Opt Shop", "whatsapp_number": "9876500001"}).json()
    assert c["whatsapp_opt_in"] is False and c["opt_in_at"] is None
    c = client.patch(f"{API}/customers/{c['id']}", headers=owner,
                     json={"whatsapp_opt_in": True, "opt_in_source": "verbal"}).json()
    assert c["whatsapp_opt_in"] is True and c["opt_in_at"] and c["opt_in_source"] == "verbal"


def test_filters(client, company_a):
    owner = company_a.h(Role.owner)
    client.post(f"{API}/customers", headers=owner, json={
        "business_name": "Dist One", "whatsapp_number": "9876500001", "customer_type": "distributor",
        "tags": ["erode"], "city": "Erode"})
    client.post(f"{API}/customers", headers=owner, json={
        "business_name": "Retail Two", "whatsapp_number": "9876500002", "tags": ["salem"]})
    assert client.get(f"{API}/customers?customer_type=distributor", headers=owner).json()["total"] == 1
    assert client.get(f"{API}/customers?tag=salem", headers=owner).json()["total"] == 1
    assert client.get(f"{API}/customers?city=erode", headers=owner).json()["total"] == 1


def test_validation_messages(client, company_a):
    r = client.post(f"{API}/customers", headers=company_a.h(Role.owner),
                    json={"business_name": "X Shop", "whatsapp_number": "9876500001", "pincode": "12"})
    assert r.status_code == 422
